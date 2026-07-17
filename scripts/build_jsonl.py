#!/usr/bin/env python3
"""把手語標記表轉成語言模型訓練用的 JSONL。

輸入（標記表）：
  1. SLR_句子標記表_台灣地名應用句子.xlsx（P01–P05，含 NMS 欄位）
  2. SLR_V2_標記表_已修改.xlsx（S01–S30 常用語句）
  3. 台灣地名與區域詞彙_手語影片命名規則.docx（W01–W38 詞彙清單，表4）

輸出（data/）：
  - tsl_sentences.jsonl  句級平行資料（中文 ↔ Gloss ↔ NMS）
  - tsl_words.jsonl      詞彙級資料（地名詞彙）
  - tsl_gloss_vocab.json Gloss 詞彙總表（下游動作庫檢索的合法詞彙清單）

對應文件：臺灣手語翻譯語言模型_微調訓練計畫.md 第 3.2 節。
"""
import json
import re
from collections import OrderedDict
from pathlib import Path

import docx
import openpyxl

BASE = Path(__file__).resolve().parent.parent
OUT_DIR = BASE / "data"

P_XLSX = Path("/Users/leo/Documents/手語影片/outputs/SLR_句子標記表_台灣地名應用句子.xlsx")
S_XLSX = Path("/Users/leo/Documents/手語影片/outputs/slr_v2_marker/SLR_V2_標記表_已修改.xlsx")
W_DOCX = Path("/Users/leo/Documents/手語影片/outputs/台灣地名與區域詞彙_手語影片命名規則.docx")

PLACEHOLDER = "＿"  # 全形底線＝填空模板句（如 S24 我叫＿）


def extract_sentences(path, prefix, has_nms=False):
    """從標記總表抽出句級資料。

    標記總表一列一個 Gloss，靠「影片檔名」欄開新影片區塊；
    同一句會出現在多部影片，Gloss 序列必須完全一致，否則報錯。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["標記總表"]
    blocks, block = [], None
    for r in ws.iter_rows(min_row=3, values_only=True):
        句號, 句子, gloss, 序, _拍攝者, 檔名 = r[0], r[1], r[2], r[3], r[4], r[5]
        nms = r[11] if has_nms and len(r) > 11 else None
        if 檔名 is not None:
            if block:
                blocks.append(block)
            block = {"no": 句號, "sent": 句子, "file": 檔名, "glosses": []}
        if block is None:
            continue
        block["glosses"].append((序, gloss, nms))
    if block:
        blocks.append(block)

    sents = OrderedDict()
    for b in blocks:
        seq = [g for _, g, _ in sorted(b["glosses"], key=lambda x: x[0])]
        nmss = [n for _, _, n in b["glosses"] if n]
        sid = f"{prefix}{b['no']:02d}"
        if sid not in sents:
            sents[sid] = {"chinese": b["sent"], "gloss": seq,
                          "nms": sorted(set(nmss)), "n_videos": 1}
        else:
            s = sents[sid]
            s["n_videos"] += 1
            s["nms"] = sorted(set(s["nms"]) | set(nmss))
            if s["gloss"] != seq:
                raise ValueError(
                    f"{sid} Gloss 不一致: {s['gloss']} vs {seq} ({b['file']})")
    return sents


def extract_words(path):
    """從命名規則 docx 的詞彙清單表（欄位：編號/代碼/詞彙內容/範例檔名）抽出 W 詞彙。"""
    d = docx.Document(path)
    for t in d.tables:
        header = [c.text.strip() for c in t.rows[0].cells]
        if "代碼" in header and "詞彙內容" in header:
            i_code, i_word = header.index("代碼"), header.index("詞彙內容")
            words = OrderedDict()
            for r in t.rows[1:]:
                code = r.cells[i_code].text.strip()
                word = r.cells[i_word].text.strip()
                if re.fullmatch(r"W\d{2}", code) and word:
                    words[code] = word
            return words
    raise ValueError("找不到含「代碼/詞彙內容」欄位的詞彙清單表")


def main():
    OUT_DIR.mkdir(exist_ok=True)

    p_sents = extract_sentences(P_XLSX, "P", has_nms=True)
    s_sents = extract_sentences(S_XLSX, "S", has_nms=False)
    words = extract_words(W_DOCX)
    assert len(p_sents) == 5, f"P 句數異常: {len(p_sents)}"
    assert len(s_sents) == 30, f"S 句數異常: {len(s_sents)}"
    assert len(words) == 38, f"W 詞數異常: {len(words)}"

    sent_path = OUT_DIR / "tsl_sentences.jsonl"
    with sent_path.open("w", encoding="utf-8") as f:
        for batch, src, sents in [
            ("台灣地名應用句子", P_XLSX.name, p_sents),
            ("SLR_V2 常用語句", S_XLSX.name, s_sents),
        ]:
            for sid, v in sents.items():
                entry = {
                    "id": sid,
                    "type": "sentence",
                    "chinese": v["chinese"],
                    "gloss": v["gloss"],
                    "gloss_text": "/".join(v["gloss"]),
                    "nms": v["nms"][0] if len(v["nms"]) == 1 else (v["nms"] or None),
                    "is_template": PLACEHOLDER in v["chinese"],
                    "n_videos": v["n_videos"],
                    "batch": batch,
                    "source_file": src,
                }
                f.write(json.dumps(entry, ensure_ascii=False) + "\n")

    word_path = OUT_DIR / "tsl_words.jsonl"
    with word_path.open("w", encoding="utf-8") as f:
        for code, word in words.items():
            entry = {
                "id": code,
                "type": "word",
                "chinese": word,
                "gloss": [word],
                "gloss_text": word,
                "nms": None,
                "is_template": False,
                "batch": "台灣地名詞彙",
                "source_file": W_DOCX.name,
            }
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")

    # Gloss 詞彙總表：模型輸出的每個 Gloss 都必須在此表內（下游動作庫才有對應動作）
    vocab = OrderedDict()
    for code, word in words.items():
        vocab.setdefault(word, {"sources": []})["sources"].append(code)
    for sents in (p_sents, s_sents):
        for sid, v in sents.items():
            for g in v["gloss"]:
                if PLACEHOLDER in g:  # 佔位符不是真 Gloss，動作庫沒有對應動作
                    continue
                vocab.setdefault(g, {"sources": []})["sources"].append(sid)
    vocab_path = OUT_DIR / "tsl_gloss_vocab.json"
    vocab_path.write_text(
        json.dumps(
            {"description": "Gloss 詞彙總表（自有資料集內出現過的所有 Gloss；"
                            "sources = 出現於哪些詞彙代碼/句子代碼）",
             "count": len(vocab),
             "glosses": {g: sorted(set(v["sources"])) for g, v in vocab.items()}},
            ensure_ascii=False, indent=2),
        encoding="utf-8")

    n_sent = len(p_sents) + len(s_sents)
    print(f"OK: {sent_path.name} {n_sent} 句 / {word_path.name} {len(words)} 詞 / "
          f"{vocab_path.name} {len(vocab)} 個 Gloss")
    templates = [sid for sents in (p_sents, s_sents) for sid, v in sents.items()
                 if PLACEHOLDER in v["chinese"]]
    if templates:
        print(f"模板句（含 {PLACEHOLDER} 佔位符，訓練前需展開或排除）: {templates}")


if __name__ == "__main__":
    main()
